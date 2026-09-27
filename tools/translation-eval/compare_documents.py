from __future__ import annotations

import argparse
from collections import Counter
from difflib import SequenceMatcher
import json
import re
from pathlib import Path

from docx import Document


NUMERIC_TOKEN_PATTERN = re.compile(
    r"(?<=第)[零〇一二三四五六七八九十百千]+(?=[章节条款部分])"
    r"|[零〇一二两三四五六七八九十百千]+(?=[个支份次管])"
    r"|\b(?:zero|one|two|three|four|five|six|seven|eight|nine|ten)\b"
    r"(?=\s+(?:(?:reference|sample|test|control)\s+)?(?:tubes?|copies?|samples?|tests?|replicates?))"
    r"|\b(?:first|second|third|fourth|fifth|sixth|seventh|eighth|ninth|tenth)\b"
    r"(?=\s+(?:(?:reference|sample|test|control)\s+)?(?:tubes?|copies?|samples?|tests?|replicates?))"
    r"|\d+(?:[.,]\d+)*(?:\s*(?:%|°C|℃|mg|mL|ml|μg|µg|ug|mcg|nm|g/min|g|U|h|min|s)(?![A-Za-z]))?",
    re.IGNORECASE,
)
IDENTIFIER_PATTERN = re.compile(
    r"(?<![A-Za-z0-9])(?=[A-Z0-9-]*[A-Z])(?=[A-Z0-9-]*\d)[A-Z0-9]+(?:-[A-Z0-9]+)+(?![A-Za-z0-9])"
)
WORD_PATTERN = re.compile(r"[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)*")
LATIN_SENTENCE_PATTERN = re.compile(r"(?:\b[A-Za-z]{2,}\b[ \t,;:'\"()./\-]*){6,}")


def _document_text(document: Document) -> str:
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    for section in document.sections:
        for container in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            values.extend(paragraph.text for paragraph in container.paragraphs)
    return "\n".join(values)


def _table_dimensions(table) -> dict[str, int]:
    try:
        columns = len(table.columns)
    except Exception:
        columns = 0
        for row in table._tbl.tr_lst:
            row_columns = 0
            for cell in row.tc_lst:
                grid_span = cell.tcPr.gridSpan if cell.tcPr is not None else None
                row_columns += int(grid_span.val) if grid_span is not None else 1
            columns = max(columns, row_columns)
    return {"rows": len(table.rows), "columns": columns}


def _profile(path: Path) -> dict[str, object]:
    document = Document(path)
    text = _document_text(document)
    return {
        "path": str(path.resolve()),
        "text": text,
        "textCharCount": len(text),
        "cjkCharCount": len(re.findall(r"[\u3400-\u9fff]", text)),
        "latinSentenceResidueCount": len(LATIN_SENTENCE_PATTERN.findall(text)),
        "paragraphCount": len(document.paragraphs),
        "nonEmptyParagraphCount": sum(1 for paragraph in document.paragraphs if paragraph.text.strip()),
        "tableCount": len(document.tables),
        "tableDimensions": [_table_dimensions(table) for table in document.tables],
        "sectionCount": len(document.sections),
        "numericTokens": [match.group(0).strip() for match in NUMERIC_TOKEN_PATTERN.finditer(text)],
        "identifiers": IDENTIFIER_PATTERN.findall(text),
    }


def _normalize_numeric(value: str) -> str:
    english_numbers = {
        "zero": 0, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5,
        "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10,
    }
    english_ordinals = {
        "first": 1, "second": 2, "third": 3, "fourth": 4, "fifth": 5,
        "sixth": 6, "seventh": 7, "eighth": 8, "ninth": 9, "tenth": 10,
    }
    if value.casefold() in english_numbers:
        return str(english_numbers[value.casefold()])
    if value.casefold() in english_ordinals:
        return str(english_ordinals[value.casefold()])
    if re.fullmatch(r"[零〇一二两三四五六七八九十百千]+", value):
        digits = {"零": 0, "〇": 0, "一": 1, "二": 2, "两": 2, "三": 3, "四": 4,
                  "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
        units = {"十": 10, "百": 100, "千": 1000}
        total = 0
        pending = 0
        for character in value:
            if character in digits:
                pending = digits[character]
            else:
                total += (pending or 1) * units[character]
                pending = 0
        return str(total + pending)
    return value.casefold().replace(" ", "").replace("µ", "μ").replace("℃", "°c")


def _preservation(
    source_values: list[str],
    candidate_values: list[str],
    normalizer=lambda value: value.casefold(),
) -> dict[str, object]:
    source = Counter(normalizer(value) for value in source_values)
    candidate = Counter(normalizer(value) for value in candidate_values)
    source_display = {normalizer(value): value for value in source_values}
    candidate_display = {normalizer(value): value for value in candidate_values}
    matched = source & candidate
    missing_counter = source - candidate
    extra_counter = candidate - source
    source_count = sum(source.values())
    candidate_count = sum(candidate.values())
    matched_count = sum(matched.values())
    return {
        "sourceCount": source_count,
        "candidateCount": candidate_count,
        "matchedCount": matched_count,
        "recall": round(matched_count / source_count, 4) if source_count else 1.0,
        "precision": round(matched_count / candidate_count, 4) if candidate_count else (1.0 if not source_count else 0.0),
        "missing": sorted(
            source_display[value]
            for value, count in missing_counter.items()
            for _ in range(count)
        ),
        "extra": sorted(
            candidate_display[value]
            for value, count in extra_counter.items()
            for _ in range(count)
        ),
    }


def _word_similarity(candidate_text: str, reference_text: str) -> dict[str, float]:
    candidate_words = [word.casefold() for word in WORD_PATTERN.findall(candidate_text)]
    reference_words = [word.casefold() for word in WORD_PATTERN.findall(reference_text)]
    candidate_set = set(candidate_words)
    reference_set = set(reference_words)
    union = candidate_set | reference_set
    jaccard = len(candidate_set & reference_set) / len(union) if union else 1.0
    sequence = SequenceMatcher(None, " ".join(candidate_words), " ".join(reference_words)).ratio()
    return {"wordJaccard": round(jaccard, 4), "normalizedSequenceRatio": round(sequence, 4)}


def _public_profile(profile: dict[str, object]) -> dict[str, object]:
    return {key: value for key, value in profile.items() if key not in {"text", "numericTokens", "identifiers"}}


def _glossary_preservation(
    source_text: str,
    candidate_text: str,
    glossary: dict[str, str] | None,
) -> dict[str, object]:
    required = [
        (source_term, target_term)
        for source_term, target_term in (glossary or {}).items()
        if source_term and target_term and source_term.casefold() in source_text.casefold()
    ]
    missing = [
        f"{source_term}→{target_term}"
        for source_term, target_term in required
        if target_term.casefold() not in candidate_text.casefold()
    ]
    return {
        "requiredCount": len(required),
        "matchedCount": len(required) - len(missing),
        "missing": sorted(missing),
    }


def compare_docx(
    source_path: Path,
    candidate_path: Path,
    reference_path: Path | None = None,
    glossary: dict[str, str] | None = None,
) -> dict[str, object]:
    source = _profile(source_path)
    candidate = _profile(candidate_path)
    result: dict[str, object] = {
        "schemaVersion": 1,
        "source": _public_profile(source),
        "candidate": _public_profile(candidate),
        "numericPreservation": _preservation(
            source["numericTokens"], candidate["numericTokens"], _normalize_numeric
        ),
        "identifierPreservation": _preservation(source["identifiers"], candidate["identifiers"]),
        "glossaryPreservation": _glossary_preservation(source["text"], candidate["text"], glossary),
        "structure": {
            "paragraphCountMatch": source["paragraphCount"] == candidate["paragraphCount"],
            "tableCountMatch": source["tableCount"] == candidate["tableCount"],
            "tableDimensionsMatch": source["tableDimensions"] == candidate["tableDimensions"],
            "sectionCountMatch": source["sectionCount"] == candidate["sectionCount"],
        },
    }
    if reference_path is not None:
        reference = _profile(reference_path)
        result["reference"] = _public_profile(reference)
        result["referenceSimilarity"] = _word_similarity(candidate["text"], reference["text"])
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Compare a translated DOCX with its source and optional reference.")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--candidate", required=True, type=Path)
    parser.add_argument("--reference", type=Path)
    parser.add_argument("--glossary", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    paths = [args.source, args.candidate] + ([args.reference] if args.reference else [])
    missing = [str(path) for path in paths if path is not None and not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing inputs: {missing}")
    glossary = None
    if args.glossary:
        glossary = json.loads(args.glossary.read_text(encoding="utf-8"))
        if not isinstance(glossary, dict):
            raise ValueError("Glossary JSON must be an object mapping source terms to target terms")
    result = compare_docx(args.source, args.candidate, args.reference, glossary)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"COMPARISON_OUTPUT={args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
