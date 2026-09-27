from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import pdfplumber
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT


NUMERIC_TOKEN_PATTERN = re.compile(
    r"(?<![A-Za-z0-9])(?:\d+(?:[.,]\d+)*(?:\s*(?:%|°C|℃|mg|mL|ml|μg|ug|mcg|U|h|min|s))?)(?![A-Za-z0-9])",
    re.IGNORECASE,
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def text_profile(text: str) -> dict[str, object]:
    return {
        "textCharCount": len(text),
        "cjkCharCount": len(re.findall(r"[\u3400-\u9fff]", text)),
        "latinWordCount": len(re.findall(r"[A-Za-z]+(?:[-'][A-Za-z]+)*", text)),
        "numericTokens": sorted(set(match.group(0).strip() for match in NUMERIC_TOKEN_PATTERN.finditer(text))),
    }


def extract_pdf(path: Path) -> dict[str, object]:
    page_texts: list[str] = []
    page_sizes: list[dict[str, float]] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_texts.append(page.extract_text() or "")
            page_sizes.append({"width": round(float(page.width), 2), "height": round(float(page.height), 2)})
    full_text = "\n".join(page_texts)
    result: dict[str, object] = {
        "pageCount": len(page_texts),
        "pagesWithText": sum(1 for text in page_texts if text.strip()),
        "hasTextLayer": any(text.strip() for text in page_texts),
        "pageSizes": page_sizes,
    }
    result.update(text_profile(full_text))
    return result


def iter_table_text(document: Document) -> list[str]:
    values: list[str] = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                values.append(cell.text)
    return values


def count_related_images(document: Document) -> int:
    return sum(1 for relationship in document.part.rels.values() if relationship.reltype == RT.IMAGE)


def extract_docx(path: Path) -> dict[str, object]:
    document = Document(path)
    body_paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_texts = iter_table_text(document)
    header_paragraphs = [
        paragraph.text
        for section in document.sections
        for header in (section.header, section.first_page_header, section.even_page_header)
        for paragraph in header.paragraphs
        if paragraph.text.strip()
    ]
    footer_paragraphs = [
        paragraph.text
        for section in document.sections
        for footer in (section.footer, section.first_page_footer, section.even_page_footer)
        for paragraph in footer.paragraphs
        if paragraph.text.strip()
    ]
    heading_paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style is not None and paragraph.style.name.lower().startswith("heading")
    ]
    full_text = "\n".join(body_paragraphs + table_texts + header_paragraphs + footer_paragraphs)
    result: dict[str, object] = {
        "paragraphCount": len(document.paragraphs),
        "nonEmptyParagraphCount": sum(1 for text in body_paragraphs if text.strip()),
        "headingCount": len(heading_paragraphs),
        "tableCount": len(document.tables),
        "tableDimensions": [
            {"rows": len(table.rows), "columns": len(table.columns)} for table in document.tables
        ],
        "imageCount": count_related_images(document),
        "sectionCount": len(document.sections),
        "headerParagraphCount": len(header_paragraphs),
        "footerParagraphCount": len(footer_paragraphs),
    }
    result.update(text_profile(full_text))
    return result


def extract(path: Path) -> dict[str, object]:
    common: dict[str, object] = {
        "name": path.name,
        "path": str(path.resolve()),
        "extension": path.suffix.lower(),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
    }
    try:
        if path.suffix.lower() == ".pdf":
            common["documentType"] = "pdf"
            common.update(extract_pdf(path))
        elif path.suffix.lower() == ".docx":
            common["documentType"] = "docx"
            common.update(extract_docx(path))
        else:
            raise ValueError(f"Unsupported document type: {path}")
        common["status"] = "ok"
    except Exception as exc:  # keep the complete suite running while preserving the failure evidence
        common["documentType"] = path.suffix.lower().lstrip(".") or "unknown"
        common["status"] = "invalid"
        common["errorType"] = type(exc).__name__
        common["error"] = str(exc)
    return common


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract structural baselines for translation samples.")
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    missing = [str(path) for path in args.inputs if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing inputs: {missing}")
    payload = {"schemaVersion": 1, "documents": [extract(path) for path in args.inputs]}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"BASELINE_OUTPUT={args.output.resolve()}")
    print(f"DOCUMENT_COUNT={len(payload['documents'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
