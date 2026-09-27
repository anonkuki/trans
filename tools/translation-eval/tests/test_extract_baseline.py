import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas


SCRIPT = Path(__file__).resolve().parents[1] / "extract_baseline.py"


class ExtractBaselineTest(unittest.TestCase):
    def test_extracts_pdf_and_docx_structure(self):
        with tempfile.TemporaryDirectory(dir="D:/Temp") as temp_dir:
            root = Path(temp_dir)
            pdf_path = root / "sample.pdf"
            pdf_canvas = canvas.Canvas(str(pdf_path))
            pdf_canvas.drawString(72, 720, "Dose 5 mg at 20 C")
            pdf_canvas.showPage()
            pdf_canvas.save()

            docx_path = root / "sample.docx"
            document = Document()
            document.add_heading("Protocol", level=1)
            document.add_paragraph("Dose 5 mg at 20 C")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Item"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "Dose"
            table.cell(1, 1).text = "5 mg"
            document.sections[0].header.paragraphs[0].text = "Header"
            document.save(docx_path)

            output_path = root / "baseline.json"
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), "--output", str(output_path), str(pdf_path), str(docx_path)],
                text=True,
                capture_output=True,
                check=False,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            payload = json.loads(output_path.read_text(encoding="utf-8"))
            by_name = {item["name"]: item for item in payload["documents"]}

            self.assertEqual(by_name["sample.pdf"]["pageCount"], 1)
            self.assertTrue(by_name["sample.pdf"]["hasTextLayer"])
            self.assertIn("5 mg", by_name["sample.pdf"]["numericTokens"])
            self.assertEqual(by_name["sample.docx"]["paragraphCount"], 2)
            self.assertEqual(by_name["sample.docx"]["tableCount"], 1)
            self.assertEqual(by_name["sample.docx"]["sectionCount"], 1)
            self.assertEqual(by_name["sample.docx"]["headerParagraphCount"], 1)

    def test_records_invalid_docx_without_aborting_other_samples(self):
        with tempfile.TemporaryDirectory(dir="D:/Temp") as temp_dir:
            root = Path(temp_dir)
            invalid_docx = root / "invalid.docx"
            invalid_docx.write_bytes(b"PK\x03\x04truncated")
            output_path = root / "baseline.json"

            completed = subprocess.run(
                [sys.executable, str(SCRIPT), "--output", str(output_path), str(invalid_docx)],
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            payload = json.loads(output_path.read_text(encoding="utf-8"))
            invalid = payload["documents"][0]
            self.assertEqual(invalid["status"], "invalid")
            self.assertEqual(invalid["errorType"], "BadZipFile")


if __name__ == "__main__":
    unittest.main()
