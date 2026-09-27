from pathlib import Path
import sys
import zipfile

from docx import Document


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from compare_documents import compare_docx  # noqa: E402


def _write_docx(path: Path, paragraphs: list[str], table_values: list[list[str]]) -> None:
    document = Document()
    for value in paragraphs:
        document.add_paragraph(value)
    if table_values:
        table = document.add_table(rows=len(table_values), cols=len(table_values[0]))
        for row_index, values in enumerate(table_values):
            for column_index, value in enumerate(values):
                table.cell(row_index, column_index).text = value
    document.save(path)


def test_compare_docx_reports_preservation_and_structure(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    reference = tmp_path / "reference.docx"
    _write_docx(source, ["温度 20 ℃，文件 BJ-SOP-TM-0201-05。"], [["1", "2"]])
    _write_docx(candidate, ["Temperature 20 ℃, document BJ-SOP-TM-0201-05."], [["1", "2"]])
    _write_docx(reference, ["Temperature: 20 ℃. Document BJ-SOP-TM-0201-05."], [["1", "2"]])

    result = compare_docx(source, candidate, reference)

    assert result["numericPreservation"]["recall"] == 1.0
    assert result["identifierPreservation"]["recall"] == 1.0
    assert result["structure"]["tableCountMatch"] is True
    assert result["structure"]["tableDimensionsMatch"] is True
    assert result["candidate"]["cjkCharCount"] == 0
    assert result["referenceSimilarity"]["wordJaccard"] > 0.5


def test_compare_docx_exposes_missing_numbers_and_table_shape(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["取样 10 mL，重复 3 次。"], [["A", "B"]])
    _write_docx(candidate, ["Take a sample and repeat 3 times."], [["A"]])

    result = compare_docx(source, candidate)

    assert result["numericPreservation"]["recall"] < 1.0
    assert "10 mL" in result["numericPreservation"]["missing"]
    assert result["structure"]["tableDimensionsMatch"] is False
    assert "referenceSimilarity" not in result


def test_numeric_preservation_normalizes_equivalent_unit_notation(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["20℃；0.1ml；100µg"], [])
    _write_docx(candidate, ["20 °C; 0.1 mL; 100 μg"], [])

    result = compare_docx(source, candidate)

    assert result["numericPreservation"]["recall"] == 1.0
    assert result["numericPreservation"]["precision"] == 1.0


def test_numeric_preservation_handles_numbers_adjacent_to_labels(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["pH10；750nm；5.5.3提前；R2"], [])
    _write_docx(candidate, ["pH 10; 750 nm; 5.5.3Advance; R2"], [])

    result = compare_docx(source, candidate)

    assert result["numericPreservation"]["recall"] == 1.0
    assert result["numericPreservation"]["precision"] == 1.0


def test_numeric_preservation_normalizes_chinese_section_ordinals(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["第二章 一般信息"], [])
    _write_docx(candidate, ["Chapter 2 General Information"], [])

    result = compare_docx(source, candidate)

    assert result["numericPreservation"]["recall"] == 1.0
    assert result["numericPreservation"]["precision"] == 1.0


def test_numeric_preservation_normalizes_spelled_out_sample_counts(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["Two reference tubes were prepared in parallel."], [])
    _write_docx(candidate, ["平行制备2支标准管。"], [])

    result = compare_docx(source, candidate)

    assert result["numericPreservation"]["recall"] == 1.0
    assert result["numericPreservation"]["precision"] == 1.0


def test_numeric_preservation_normalizes_chinese_spelled_out_counts(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["Two reference tubes were prepared in parallel."], [])
    _write_docx(candidate, ["平行制备两个参考管。"], [])

    result = compare_docx(source, candidate)

    assert result["numericPreservation"]["recall"] == 1.0
    assert result["numericPreservation"]["precision"] == 1.0


def test_numeric_preservation_normalizes_english_sample_ordinals(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["Start timing when the first tube is dosed."], [])
    _write_docx(candidate, ["从第一管加样时开始计时。"], [])

    result = compare_docx(source, candidate)

    assert result["numericPreservation"]["recall"] == 1.0
    assert result["numericPreservation"]["precision"] == 1.0


def test_identifier_preservation_handles_cjk_adjacency(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["依据CC-2018-0137修订"], [])
    _write_docx(candidate, ["Revised according to CC-2018-0137"], [])

    result = compare_docx(source, candidate)

    assert result["identifierPreservation"]["sourceCount"] == 1
    assert result["identifierPreservation"]["recall"] == 1.0
    assert result["identifierPreservation"]["precision"] == 1.0


def test_compare_docx_reports_required_glossary_targets(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["执行污染控制策略并调查超标结果。"], [])
    _write_docx(candidate, ["Implement the contamination strategy and investigate excursions."], [])

    result = compare_docx(
        source,
        candidate,
        glossary={"污染控制策略": "contamination control strategy", "超标": "excursion"},
    )

    assert result["glossaryPreservation"]["requiredCount"] == 2
    assert result["glossaryPreservation"]["matchedCount"] == 1
    assert result["glossaryPreservation"]["missing"] == ["污染控制策略→contamination control strategy"]


def test_compare_docx_handles_legacy_table_without_tbl_grid(tmp_path: Path) -> None:
    source = tmp_path / "legacy-gridless.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["表格"], [["A", "B"], ["C", "D"]])
    _write_docx(candidate, ["Table"], [["A", "B"], ["C", "D"]])

    rewritten = tmp_path / "rewritten.docx"
    with zipfile.ZipFile(source) as input_zip, zipfile.ZipFile(rewritten, "w") as output_zip:
        for item in input_zip.infolist():
            content = input_zip.read(item.filename)
            if item.filename == "word/document.xml":
                start = content.index(b"<w:tblGrid>")
                end = content.index(b"</w:tblGrid>", start) + len(b"</w:tblGrid>")
                content = content[:start] + content[end:]
            output_zip.writestr(item, content)
    rewritten.replace(source)

    result = compare_docx(source, candidate)

    assert result["source"]["tableDimensions"] == [{"rows": 2, "columns": 2}]
    assert result["structure"]["tableDimensionsMatch"] is True


def test_compare_docx_counts_long_latin_sentence_residue(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    candidate = tmp_path / "candidate.docx"
    _write_docx(source, ["制备测试样品管。"], [])
    _write_docx(candidate, ["According to the source of the test product, prepare the test sample tube.依"], [])

    result = compare_docx(source, candidate)

    assert result["candidate"]["latinSentenceResidueCount"] == 1
